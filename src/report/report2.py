from pptx import Presentation
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.util import Inches, Pt
from pptx import Presentation
from pptx.chart.data import ChartData, CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_TICK_MARK, XL_TICK_LABEL_POSITION, XL_MARKER_STYLE, XL_LABEL_POSITION
from pptx.util import Cm  # Inches
from pandas import np
from pptx.enum.chart import XL_LEGEND_POSITION
import math
from collections import Counter, OrderedDict
from report_common import ReportBase
import report_common
import report_logger

from pptx.dml.color import RGBColor
import pandas as pd

class Report2(ReportBase):

    def Run(self):
        df_vw = self.df_vw
        df_mkt = self.df_mkt

        PR_Status_actual = self.PR_Status_actual
        PR_Status_local = self.PR_Status_local
        PR_Status = self.PR_Status
        year_filter = self.year_filter
        oem_filter = self.oem_filter
        brand_filter = self.brand_filter
        build_type_filter = self.build_type_filter
        shapes = self.shapes
        year_categories = self.year_categories
        view_rules = self.view_rules  # Brand or OEM
        all_mkt = self.all_mkt  # true分母是否为全市场，false分母为fuel_type细分市场

        fuel_type_group_filter = self.fuel_type_group_filter
        fuel_type_group_all = self.fuel_type_group_all
        fuel_type_filter = self.fuel_type_filter
        fuel_type_all = self.fuel_type_all


        # 根据所有filter过滤大众数据
        df_vw_filter = self.df_vw_filter

        # 根据所有filter过滤mkt数据
        df_mkt_filter = self.df_mkt_filter

        # 全销售市场数据每year每fueltype分组聚合(ice bev phev)-------------------------
        df_mkt_group = df_mkt_filter.groupby(['Fuel_type', 'Year']) \
            .agg({'Volume': np.sum}).reset_index()
        df_mkt_group_premium = df_mkt_filter[(df_mkt_filter['Brand_Indicator'] == 'Premium')].groupby(['Fuel_type', 'Year'])\
            .agg({'Volume': np.sum}).reset_index()

        # 按大众本轮的销量排序OEM和Brand
        categories_order = \
            df_vw_filter[df_vw_filter['PR_Status'].isin(PR_Status)].groupby(view_rules).agg({'Volume': np.sum}) \
            .reset_index().sort_values(['Volume'], ascending=[False]).reset_index()[view_rules]
        # print(categories_order)

        # 大众市场销量数据分组统计
        df_vw_categories_year = df_vw_filter.groupby([view_rules, 'YEAR']).agg({'Volume': np.sum}).reset_index()
        # print(df_vw_categories_year)

        # 大众市场奥迪豪华车销量数据分组统计
        df_vw_audi_premium_year = \
            df_vw_filter[(df_vw_filter['Brand'] == 'Audi') & (df_vw_filter['Brand_Indicator'] == 'Premium')] \
            .groupby(['YEAR']).agg({'Volume': np.sum}).reset_index()
        # print(df_vw_audi_premium_year)

        if all_mkt:
            calculate_fuel_type = fuel_type_all
        else:
            calculate_fuel_type = fuel_type_filter

        category_volume_dict = OrderedDict()
        category_mkt_rate_dict = OrderedDict()
        category_totat_mkt_rate_dict = OrderedDict()
        for category in categories_order:
            df_category_volume = df_vw_categories_year[(df_vw_categories_year[view_rules] == category)].reset_index()
            category_year_vol_dict = OrderedDict()
            category_year_rate_dict = OrderedDict()
            total_mke_rate = 0
            for year in year_filter:
                year_volume = 0
                mkt_rate = 0
                if not df_category_volume[(df_category_volume['YEAR'] == year)].empty:
                    year_volume = df_category_volume[(df_category_volume['YEAR'] == year)].reset_index().at[0, 'Volume']
                    mkt_volume_sum = 0
                    for fuel in calculate_fuel_type:
                        mkt_volume = 0
                        if not df_mkt_group[(df_mkt_group['Year'] == year) & (df_mkt_group['Fuel_type'] == fuel)].empty:
                            mkt_volume = \
                                df_mkt_group[(df_mkt_group['Year'] == year) & (df_mkt_group['Fuel_type'] == fuel)] \
                                .reset_index().at[0, 'Volume']
                        else:
                            report_logger.record_log('report' + str(self.id), 'VW', PR_Status_local, year, fuel, '', category)
                        mkt_volume_sum = mkt_volume_sum + mkt_volume
                    if mkt_volume_sum > 0:
                        mkt_rate = year_volume / mkt_volume_sum
                    total_mke_rate = total_mke_rate + mkt_rate
                else:
                    report_logger.record_log('report'+str(self.id), 'VW', PR_Status_local, year, '', '', category)
                category_year_vol_dict[year] = year_volume
                category_year_rate_dict[year] = mkt_rate
            category_volume_dict[category] = category_year_vol_dict
            category_mkt_rate_dict[category] = category_year_rate_dict
            category_totat_mkt_rate_dict[category] = total_mke_rate

        #单独计算奥迪豪华市场的市占率
        year_rate_audi_premium_dict = OrderedDict()
        for year in year_filter:
            mkt_rate = 0
            year_volume = 0
            if not df_vw_audi_premium_year[(df_vw_audi_premium_year['YEAR'] == year)].empty:
                year_volume = \
                    df_vw_audi_premium_year[(df_vw_audi_premium_year['YEAR'] == year)].reset_index().at[0, 'Volume']
                mkt_volume_sum = 0
                for fuel in calculate_fuel_type:
                    mkt_volume = 0
                    if not df_mkt_group_premium[(df_mkt_group_premium['Year'] == year) & (df_mkt_group_premium['Fuel_type'] == fuel)].empty:
                        mkt_volume = \
                            df_mkt_group_premium[(df_mkt_group_premium['Year'] == year) & (df_mkt_group_premium['Fuel_type'] == fuel)] \
                            .reset_index().at[0, 'Volume']
                    else:
                        report_logger.record_log('report' + str(self.id), 'VW', PR_Status_local, year, fuel, '', 'Audi')
                    mkt_volume_sum = mkt_volume_sum + mkt_volume
                if mkt_volume_sum > 0:
                    mkt_rate = year_volume / mkt_volume_sum
            else:
                report_logger.record_log('report' + str(self.id), 'VW', PR_Status_local, year, '', '', 'Audi')
            year_rate_audi_premium_dict[year] = mkt_rate


        top_base = 2
        # 开始创建点线图-----------------------------------
        chart_data_line = ChartData()
        chart_data_line.categories = year_categories
        for category in categories_order:
            chart_data_line.add_series(category, category_mkt_rate_dict[category].values())

        x, y, cx, cy = Cm(1), Cm(top_base), Cm(24), Cm(7)
        chart_line = shapes.add_chart(
            XL_CHART_TYPE.LINE, x, y, cx, cy, chart_data_line
        ).chart


        chart_line.has_legend = True
        chart_line.legend.include_in_layout = False
        chart_line.legend.position = XL_LEGEND_POSITION.LEFT
        chart_line.legend.font.size = Pt(10)

        for line_serie in chart_line.series:
            line_serie.smooth = True
            line_serie.marker.style = XL_MARKER_STYLE.CIRCLE
            line_serie.data_labels.show_value = True
            line_serie.data_labels.number_format = '0.0%'
            line_serie.data_labels.font.size = Pt(8)
            line_serie.data_labels.position = XL_LABEL_POSITION.RIGHT #XL_LABEL_POSITION.ABOVE


        value_axis_line = chart_line.value_axis
        value_axis_line.has_major_gridlines = False
        value_axis_line.major_tick_mark = XL_TICK_MARK.NONE
        value_axis_line.tick_label_position = XL_TICK_LABEL_POSITION.NONE
        value_axis_line.format.line.dash_style = MSO_LINE_DASH_STYLE.ROUND_DOT
        value_axis_line.visible = False

        category_axis_line = chart_line.category_axis
        category_axis_line.has_major_gridlines = False
        category_axis_line.major_tick_mark = XL_TICK_MARK.NONE
        category_axis_line.tick_label_position = XL_TICK_LABEL_POSITION.NONE
        category_axis_line.format.line.dash_style = MSO_LINE_DASH_STYLE.ROUND_DOT
        category_axis_line.visible = False

        # 开始创建表格
        if view_rules == 'Brand':
            rows = len(categories_order) * 2 + 2
        else:
            rows = len(categories_order) * 2 + 1

        cols = len(year_categories) + 1
        table_width = 24
        table_height = 3
        top = Cm(top_base + 7)
        left = Cm(1.5)  # Inches(2.0)
        width = Cm(table_width)  # Inches(6.0)
        height = Cm(table_height)  # Inches(0.8)

        # 添加表格到幻灯片 --------------------
        table = shapes.add_table(rows, cols, left, top, width, height).table

        # 设置单元格宽度
        columns_width = table_width / cols - 0.1
        for i in range(cols):
            if i == 0:
                table.columns[i].width = Cm(columns_width + 0.6)
            else:
                table.columns[i].width = Cm(columns_width)  # Inches(2.0)

        row_height = table_height / rows
        for i in range(rows):
            table.rows[i].height = Cm(row_height)  # Inches(2.0)

        # 设置标题行
        for i in range(cols):
            if i == 0:
                table.cell(0, i).text = 'Vol.by ' + view_rules
            else:
                table.cell(0, i).text = str(year_categories[i - 1])


        #填充每行的数据
        category_idx = 0
        last_row_idx = 0
        for idx, category in enumerate(categories_order):
            row_idx = idx * 2
            table.cell(row_idx + 1, 0).text = category
            table.cell(row_idx + 2, 0).text = 'MKT%'
            for col_idx, year in enumerate(year_filter):
                volume = category_year_vol_dict[year]
                rate = category_year_rate_dict[year]
                table.cell(row_idx + 1, col_idx + 1).text = format(volume / 1000, '.0f')
                table.cell(row_idx + 2, col_idx + 1).text = format(rate, '.1%')
            if category == 'Audi':
                table.cell(row_idx + 3, 0).text = 'Premium%'
                for col_idx, year in enumerate(year_filter):
                    rate = year_rate_audi_premium_dict[year]
                    table.cell(row_idx + 3, col_idx + 1).text = format(rate, '.1%')
                last_row_idx = row_idx + 3 + 1
                category_idx = idx
                break

        for idx, category in enumerate(categories_order):
            if idx > category_idx:
                row_idx = last_row_idx + (idx - category_idx - 1) * 2
                table.cell(row_idx, 0).text = category
                table.cell(row_idx + 1, 0).text = 'MKT%'
                for col_idx, year in enumerate(year_filter):
                    volume = category_year_vol_dict[year]
                    rate = category_year_rate_dict[year]
                    table.cell(row_idx, col_idx + 1).text = format(volume / 1000, '.0f')
                    table.cell(row_idx + 1, col_idx + 1).text = format(rate, '.1%')

        # 调整table每个cell的字体
        report_common.set_table_format(table, Pt(8), r'/')

        # 开始添加注释文本框
        left = Cm(1)  # left，top为相对位置
        top = Cm(top_base - 0.8)
        width = Cm(2)  # width，height为文本框的大小
        height = Cm(1)

        # 在指定位置添加文本框
        textbox = shapes.add_textbox(left, top, width, height)
        report_common.set_textbox_format(textbox, Pt(8), "Volume\n'000units")


